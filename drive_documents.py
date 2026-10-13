#!/usr/bin/env python3
"""Read supported documents from an already-scanned Google Drive folder."""

from __future__ import annotations

import argparse
from hashlib import sha256
from html.parser import HTMLParser
from io import BytesIO
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from xml.etree import ElementTree

from drive_scan import google_service
from output_layout import is_promotional_document, output_paths, safe_output_file


MAX_DOCUMENT_BYTES = 10_000_000
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_MIME = "text/plain"
MARKDOWN_MIME = "text/markdown"
WORD_MIME = "application/msword"
GOOGLE_DOC_MIME = "application/vnd.google-apps.document"
LEGACY_PPT_MIME = "application/vnd.ms-powerpoint"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
HTML_MIME = "text/html"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSM_MIME = "application/vnd.ms-excel.sheet.macroEnabled.12"
XLS_MIME = "application/vnd.ms-excel"
GOOGLE_SHEET_MIME = "application/vnd.google-apps.spreadsheet"


class DocumentHTMLParser(HTMLParser):
    """Keep visible text from HTML files stored with a .doc extension."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        elif tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)
        elif tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def decode_document_text(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("utf-16")


def effective_document_mime(name: str, mime_type: str) -> str:
    """Use a known extension when Drive labels a document as generic binary."""
    if mime_type != "application/octetstream":
        return mime_type
    return {".pdf": PDF_MIME, ".docx": DOCX_MIME, ".pptx": PPTX_MIME,
            ".doc": WORD_MIME, ".txt": TEXT_MIME, ".md": MARKDOWN_MIME,
            ".html": HTML_MIME, ".htm": HTML_MIME, ".xlsx": XLSX_MIME,
            ".xlsm": XLSM_MIME, ".xls": XLS_MIME}.get(
                Path(name).suffix.casefold(), mime_type)


def extract_text(raw: bytes, mime_type: str) -> str:
    if mime_type in {TEXT_MIME, MARKDOWN_MIME}:
        return decode_document_text(raw)
    if mime_type == HTML_MIME:
        parser = DocumentHTMLParser()
        parser.feed(decode_document_text(raw))
        return re.sub(r"\n[ \t]*\n[ \t]*\n+", "\n\n", "".join(parser.parts))
    if mime_type in {XLSX_MIME, XLSM_MIME}:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            sections = []
            for sheet in workbook.worksheets:
                lines = ["\t".join("" if value is None else str(value) for value in row).rstrip("\t")
                         for row in sheet.iter_rows(values_only=True)]
                sections.append(f"[Trang tính {sheet.title}]\n" + "\n".join(line for line in lines if line))
            return "\n\n".join(sections)
        finally:
            workbook.close()
    if mime_type == XLS_MIME:
        if not shutil.which("libreoffice"):
            raise RuntimeError("Thiếu LibreOffice để đọc file Excel .xls")
        with TemporaryDirectory(prefix="getsub-xls-") as folder:
            source = Path(folder) / "input.xls"
            source.write_bytes(raw)
            command = ["libreoffice", "-env:UserInstallation=" + (Path(folder) / "profile").as_uri(),
                       "--headless", "--convert-to", "xlsx", "--outdir", folder, str(source)]
            result = subprocess.run(command, capture_output=True, timeout=90, check=False)
            converted = Path(folder) / "input.xlsx"
            if result.returncode != 0 or not converted.is_file():
                raise ValueError("LibreOffice không chuyển được Excel .xls")
            return extract_text(converted.read_bytes(), XLSX_MIME)
    if mime_type == WORD_MIME:
        if raw.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            if not shutil.which("catdoc"):
                raise RuntimeError("Thiếu catdoc để đọc file Word .doc")
            with TemporaryDirectory(prefix="getsub-doc-") as folder:
                source = Path(folder) / "input.doc"
                source.write_bytes(raw)
                result = subprocess.run(["catdoc", "-w", "-d", "utf-8", str(source)],
                                        capture_output=True, timeout=45, check=False)
            if result.returncode != 0:
                raise ValueError("catdoc không đọc được file Word .doc")
            return result.stdout.decode("utf-8", errors="replace")
        if raw.lstrip().lower().startswith((b"<html", b"<!doctype html")) or raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            text = decode_document_text(raw)
            if not text.lstrip().lower().startswith(("<html", "<!doctype html")):
                raise ValueError("File .doc không có cấu trúc Word hoặc HTML")
            parser = DocumentHTMLParser()
            parser.feed(text)
            return re.sub(r"\n[ \t]*\n[ \t]*\n+", "\n\n", "".join(parser.parts))
        raise ValueError("File .doc không có cấu trúc Word hoặc HTML")
    if mime_type == LEGACY_PPT_MIME:
        if not raw.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise ValueError("File .ppt không có cấu trúc PowerPoint cũ")
        if not shutil.which("libreoffice"):
            raise RuntimeError("Thiếu LibreOffice để đọc file PowerPoint .ppt")
        with TemporaryDirectory(prefix="getsub-ppt-") as folder:
            source = Path(folder) / "input.ppt"
            source.write_bytes(raw)
            command = ["libreoffice", "-env:UserInstallation=" + (Path(folder) / "profile").as_uri(),
                       "--headless", "--convert-to", "pptx", "--outdir", folder, str(source)]
            result = subprocess.run(command, capture_output=True, timeout=90, check=False)
            converted = Path(folder) / "input.pptx"
            if result.returncode != 0 or not converted.is_file():
                raise ValueError("LibreOffice không chuyển được PowerPoint .ppt")
            return extract_text(converted.read_bytes(), PPTX_MIME)
    if mime_type == PDF_MIME:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(raw), strict=False)
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    if mime_type == DOCX_MIME:
        from docx import Document

        document = Document(BytesIO(raw))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if mime_type == PPTX_MIME:
        with ZipFile(BytesIO(raw)) as slides:
            names = sorted((name for name in slides.namelist()
                            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
                           key=lambda name: int(re.search(r"slide(\d+)", name).group(1)))
            sections = []
            for number, name in enumerate(names, 1):
                root = ElementTree.fromstring(slides.read(name))
                parts = [node.text.strip() for node in root.iter()
                         if node.tag.endswith("}t") and node.text and node.text.strip()]
                if parts:
                    sections.append(f"[Slide {number}]\n" + "\n".join(parts))
            return "\n\n".join(sections)
    raise ValueError(f"Chưa hỗ trợ trích văn bản từ MIME: {mime_type}")


def ocr_pdf(raw: bytes) -> str:
    """OCR at most 40 pages locally; also accepts image files mislabeled as PDF."""
    local = Path(__file__).resolve().parent / ".deps/ocr"
    tesseract = shutil.which("tesseract") or str(local / "usr/bin/tesseract")
    if not Path(tesseract).is_file():
        raise RuntimeError("Thiếu Tesseract; cài tesseract-ocr cùng dữ liệu vie/eng")
    tessdata = local / "usr/share/tesseract-ocr/5/tessdata"
    env = os.environ.copy()
    if tessdata.is_dir():
        env["TESSDATA_PREFIX"] = str(tessdata)
    with TemporaryDirectory(prefix="getsub-ocr-") as temp:
        root = Path(temp)
        if raw.startswith((b"\xff\xd8", b"\x89PNG")):
            image = root / ("image.jpg" if raw.startswith(b"\xff\xd8") else "image.png")
            image.write_bytes(raw)
            images = [image]
        else:
            pdf = root / "input.pdf"
            pdf.write_bytes(raw)
            command = ["pdftoppm", "-f", "1", "-l", "40", "-r", "120", "-jpeg", str(pdf), str(root / "page")]
            run = subprocess.run(command, capture_output=True, text=True, timeout=180)
            if run.returncode != 0:
                raise RuntimeError(f"Không chuyển được PDF sang ảnh: {run.stderr[-300:]}")
            images = sorted(root.glob("page-*.jpg"))
        sections = []
        for index, image in enumerate(images, 1):
            run = subprocess.run([tesseract, str(image), "stdout", "-l", "vie+eng"],
                                 capture_output=True, text=True, timeout=120, env=env)
            if run.returncode != 0:
                raise RuntimeError(f"OCR lỗi ở trang {index}: {run.stderr[-300:]}")
            if run.stdout.strip():
                sections.append(f"[Trang {index}]\n{run.stdout.strip()}")
        return "\n\n".join(sections)


def save_manifest(path: Path, scan: dict, results: list[dict]) -> dict:
    counts = {status: sum(x["status"] == status for x in results) for status in (
        "ok", "excluded", "empty", "needs_ocr", "too_large", "unsupported", "no_access", "error"
    )}
    manifest = {"folder_id": scan["folder"]["id"], "total": len(results), **counts, "files": results}
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)
    return manifest


def download_media(service, file_id: str, size: int | None, max_bytes: int) -> bytes:
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    if not size or size <= 10_000_000:
        return request.execute()
    from googleapiclient.http import MediaIoBaseDownload

    stream = BytesIO()
    downloader = MediaIoBaseDownload(stream, request, chunksize=2_000_000)
    done = False
    while not done:
        status, done = downloader.next_chunk(num_retries=2)
        if status is not None:
            print(f"Đang tải tài liệu lớn {file_id}: {status.progress():.0%}",
                  file=sys.stderr, flush=True)
        if stream.tell() > max_bytes:
            raise ValueError(f"Tệp tải về vượt giới hạn {max_bytes} byte")
    return stream.getvalue()


def process_documents(service, scan: dict, output_dir: Path, max_bytes: int = MAX_DOCUMENT_BYTES,
                      refresh: bool = False, ocr: bool = False) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output_dir / "manifest.json"
    named_paths = output_paths(scan, "document")
    previous = {}
    if manifest_path.is_file() and not refresh:
        old = json.loads(manifest_path.read_text(encoding="utf-8"))
        previous = {entry["id"]: entry for entry in old.get("files", [])}
    results = []
    for item in scan["items"]:
        if item["kind"] != "document":
            continue
        entry = {
            "id": item["id"],
            "name": item["name"],
            "path": item["path"],
            "mime_type": item["mime_type"],
            "url": item["url"],
            "modified_time": item.get("modified_time"),
        }
        try:
            if not re.fullmatch(r"[A-Za-z0-9_-]+", item["id"]):
                raise ValueError("ID file không hợp lệ")
            output_name = named_paths[item["id"]].as_posix()
            output_path = safe_output_file(output_dir, output_name, ".txt")
            prior = previous.get(item["id"])
            legacy_path = safe_output_file(output_dir, item["id"] + ".txt", ".txt")
            old_name = prior.get("output", item["id"] + ".txt") if prior else item["id"] + ".txt"
            old_path = safe_output_file(output_dir, old_name, ".txt")
            mime_type = effective_document_mime(item["name"], item["mime_type"])
            if item["name"].startswith("~$") and Path(item["name"]).suffix.casefold() in {".doc", ".docx", ".ppt", ".pptx"}:
                entry.update({"status": "excluded", "reason": "File khóa tạm của Office"})
                results.append(entry)
                if len(results) % 50 == 0:
                    save_manifest(manifest_path, scan, results)
                continue
            if is_promotional_document(item):
                if old_path.is_file() and is_promotional_document(item, old_path.read_text(encoding="utf-8")):
                    old_path.unlink()
                entry.update({"status": "excluded", "reason": "Tài liệu quảng cáo Khóa học giá hời"})
                results.append(entry)
                if len(results) % 50 == 0:
                    save_manifest(manifest_path, scan, results)
                continue
            if (not refresh and prior and prior.get("modified_time") == item.get("modified_time")
                    and prior.get("status") in ({"empty", "no_access"} if ocr else
                                                {"empty", "needs_ocr", "no_access"})):
                entry.update(prior)
                results.append(entry)
                if len(results) % 50 == 0:
                    save_manifest(manifest_path, scan, results)
                continue
            if (not refresh and prior and prior.get("status") == "ok"
                    and prior.get("modified_time") == item.get("modified_time")
                    and (old_path.is_file() or output_path.is_file())):
                cached_path = old_path if old_path.is_file() else output_path
                content = cached_path.read_text(encoding="utf-8").strip()
                if content and sha256(content.encode("utf-8")).hexdigest() == prior.get("sha256"):
                    if is_promotional_document(item, content):
                        cached_path.unlink()
                        entry.update({"status": "excluded", "reason": "Tài liệu quảng cáo Khóa học giá hời"})
                    else:
                        output_path.parent.mkdir(parents=True, exist_ok=True)
                        if cached_path != output_path:
                            cached_path.replace(output_path)
                        entry.update(prior)
                        entry["output"] = output_name
                    results.append(entry)
                    if len(results) % 50 == 0:
                        save_manifest(manifest_path, scan, results)
                    continue
            # A run interrupted before its first manifest can still reuse extracted text.
            cached_path = output_path if output_path.is_file() else legacy_path
            if not refresh and not prior and cached_path.is_file():
                content = cached_path.read_text(encoding="utf-8").strip()
                if content:
                    if is_promotional_document(item, content):
                        cached_path.unlink()
                        entry.update({"status": "excluded", "reason": "Tài liệu quảng cáo Khóa học giá hời"})
                    else:
                        output_path.parent.mkdir(parents=True, exist_ok=True)
                        if cached_path != output_path:
                            cached_path.replace(output_path)
                        entry.update({"status": "ok", "bytes": item.get("size"),
                                      "characters": len(content),
                                      "sha256": sha256(content.encode("utf-8")).hexdigest(),
                                      "output": output_name})
                    results.append(entry)
                    if len(results) % 50 == 0:
                        save_manifest(manifest_path, scan, results)
                    continue
            if item.get("can_download") is False:
                entry.update({"status": "no_access"})
            elif item.get("size") == 0 and mime_type != GOOGLE_DOC_MIME:
                entry.update({"status": "empty", "bytes": 0})
            elif item.get("size") and item["size"] > max_bytes:
                entry.update({"status": "too_large", "bytes": item["size"]})
            elif mime_type not in {PDF_MIME, DOCX_MIME, TEXT_MIME, MARKDOWN_MIME,
                                   WORD_MIME, GOOGLE_DOC_MIME, PPTX_MIME, LEGACY_PPT_MIME,
                                   HTML_MIME, XLSX_MIME, XLSM_MIME, XLS_MIME, GOOGLE_SHEET_MIME}:
                entry.update({"status": "unsupported"})
            else:
                raw = (service.files().export_media(fileId=item["id"],
                                                    mimeType="text/csv" if mime_type == GOOGLE_SHEET_MIME else TEXT_MIME).execute()
                       if mime_type in {GOOGLE_DOC_MIME, GOOGLE_SHEET_MIME} else
                       download_media(service, item["id"], item.get("size"), max_bytes))
                if len(raw) > max_bytes:
                    entry.update({"status": "too_large", "bytes": len(raw)})
                else:
                    if mime_type == PDF_MIME and raw.startswith((b"\xff\xd8", b"\x89PNG")):
                        content = ocr_pdf(raw).strip() if ocr else ""
                    else:
                        content = extract_text(raw, TEXT_MIME if mime_type in {GOOGLE_DOC_MIME, GOOGLE_SHEET_MIME}
                                               else mime_type).strip()
                        if not content and ocr and mime_type == PDF_MIME:
                            content = ocr_pdf(raw).strip()
                    if not content:
                        entry.update({"status": "needs_ocr" if mime_type == PDF_MIME else "empty",
                                      "bytes": len(raw)})
                    elif is_promotional_document(item, content):
                        entry.update({"status": "excluded", "bytes": len(raw),
                                      "reason": "Tài liệu quảng cáo Khóa học giá hời"})
                    else:
                        output_path.parent.mkdir(parents=True, exist_ok=True)
                        output_path.write_text(content + "\n", encoding="utf-8")
                        entry.update({
                            "status": "ok",
                            "bytes": len(raw),
                            "characters": len(content),
                            "sha256": sha256(content.encode("utf-8")).hexdigest(),
                            "output": output_name,
                        })
        except Exception as exc:
            entry.update({"status": "error", "error": str(exc)})
        results.append(entry)
        if len(results) % 50 == 0:
            save_manifest(manifest_path, scan, results)
    return save_manifest(manifest_path, scan, results)


def main() -> int:
    parser = argparse.ArgumentParser(description="Trích xuất văn bản PDF, DOCX, PPTX và TXT trong báo cáo quét Drive")
    parser.add_argument("--scan", type=Path, default=Path("drive-reports/scan.json"))
    parser.add_argument("--credentials", type=Path, default=Path("credentials.json"))
    parser.add_argument("--token", type=Path, default=Path("token.json"))
    parser.add_argument("--output-dir", type=Path, default=Path("drive-reports/documents"))
    parser.add_argument("--max-mb", type=int, default=10, help="Giới hạn tải từng tài liệu (MB)")
    parser.add_argument("--refresh", action="store_true", help="Tải và trích xuất lại mọi tài liệu")
    parser.add_argument("--ocr", action="store_true", help="Dùng Tesseract cục bộ cho PDF chỉ có ảnh")
    args = parser.parse_args()
    try:
        if args.max_mb < 1:
            raise ValueError("--max-mb phải lớn hơn 0")
        scan = json.loads(args.scan.read_text(encoding="utf-8"))
        service = google_service(args.credentials, args.token)
        result = process_documents(service, scan, args.output_dir,
                                   max_bytes=args.max_mb * 1_000_000, refresh=args.refresh, ocr=args.ocr)
        print(
            f"Tài liệu: {result['ok']}/{result['total']} đã đọc; "
            f"bỏ qua quảng cáo: {result['excluded']}; "
            f"cần OCR: {result['needs_ocr']}; quá lớn: {result['too_large']}; "
            f"không hỗ trợ: {result['unsupported']}; lỗi: {result['error']}"
        )
        return 0 if result["error"] == 0 else 1
    except (OSError, ValueError) as exc:
        print(f"Lỗi xử lý tài liệu: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
